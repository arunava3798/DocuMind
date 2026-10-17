"""Load and extract text from supported document formats (lightweight, portable libraries)."""
from __future__ import annotations

import os
from typing import Iterable

from langchain_community.document_loaders import (
    Docx2txtLoader,
    PyPDFLoader,
    TextLoader,
)
from langchain_core.documents import Document

SUPPORTED_EXTENSIONS = {
    ".pdf",
    ".txt",
    ".docx",
    ".xls",
    ".xlsx",
    ".pptx",
}


def _to_docs(text: str, filename: str) -> list[Document]:
    return [Document(page_content=text, metadata={"source": filename})]


def _load_xlsx(path: str) -> str:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        blocks = []
        for sheet in workbook.worksheets:
            rows = []
            for row in sheet.iter_rows(values_only=True):
                cells = [str(c) for c in row if c is not None]
                if cells:
                    rows.append("\t".join(cells))
            if rows:
                blocks.append(f"# Sheet: {sheet.title}\n" + "\n".join(rows))
        return "\n\n".join(blocks)
    finally:
        workbook.close()


def _load_xls(path: str) -> str:
    import xlrd

    book = xlrd.open_workbook(path)
    blocks = []
    for sheet in book.sheets():
        rows = []
        for r in range(sheet.nrows):
            cells = [str(sheet.cell_value(r, c)) for c in range(sheet.ncols)]
            cells = [c for c in cells if c.strip()]
            if cells:
                rows.append("\t".join(cells))
        if rows:
            blocks.append(f"# Sheet: {sheet.name}\n" + "\n".join(rows))
    return "\n\n".join(blocks)


def _load_pptx(path: str) -> str:
    from pptx import Presentation

    presentation = Presentation(path)
    blocks = []
    for i, slide in enumerate(presentation.slides, start=1):
        lines = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    line = "".join(run.text for run in paragraph.runs)
                    if line.strip():
                        lines.append(line)
        if lines:
            blocks.append(f"# Slide {i}\n" + "\n".join(lines))
    return "\n\n".join(blocks)


def load_file(path: str) -> list[Document]:
    """Load a single file into a list of LangChain Documents."""
    ext = os.path.splitext(path)[1].lower()
    filename = os.path.basename(path)

    if ext == ".pdf":
        docs = PyPDFLoader(path).load()
    elif ext == ".txt":
        docs = TextLoader(path, encoding="utf-8").load()
    elif ext == ".docx":
        docs = Docx2txtLoader(path).load()
    elif ext == ".xlsx":
        docs = _to_docs(_load_xlsx(path), filename)
    elif ext == ".xls":
        docs = _to_docs(_load_xls(path), filename)
    elif ext == ".pptx":
        docs = _to_docs(_load_pptx(path), filename)
    else:
        raise ValueError(f"Unsupported file type: {ext}")

    for doc in docs:
        doc.metadata["source"] = filename
    return docs


def load_paths(paths: Iterable[str]) -> list[Document]:
    """Load several files given their paths."""
    documents: list[Document] = []
    for path in paths:
        documents.extend(load_file(path))
    return documents
